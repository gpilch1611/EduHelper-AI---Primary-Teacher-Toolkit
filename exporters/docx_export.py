# exporters/docx_export.py – Lesson plan → DOCX via python-docx

from __future__ import annotations

import re
from pathlib import Path


def export_docx(content: str, filepath: str | Path) -> None:
    """
    Render a lesson plan text to a formatted .docx file.

    Raises ImportError if python-docx is not installed.
    """
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor, Cm
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
    except ImportError:
        raise ImportError(
            "python-docx is not installed.\n"
            "Run:  pip install python-docx"
        )

    doc = Document()

    # ── Page margins ──────────────────────────────────────────────────────────
    for section in doc.sections:
        section.top_margin    = Cm(2.0)
        section.bottom_margin = Cm(2.5)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    # ── Colour constants ──────────────────────────────────────────────────────
    BLUE  = RGBColor(0x25, 0x63, 0xEB)
    DARK  = RGBColor(0x1A, 0x1A, 0x1A)
    GREY  = RGBColor(0x55, 0x55, 0x55)

    # ── Style helpers ─────────────────────────────────────────────────────────

    def _heading(text: str, level: int = 1) -> None:
        para = doc.add_heading("", level=level)
        run  = para.add_run(text)
        run.font.color.rgb = BLUE
        run.font.bold = True
        run.font.size = Pt(14 if level == 1 else 12)

    def _phase_heading(text: str) -> None:
        para = doc.add_paragraph()
        para.paragraph_format.space_before = Pt(8)
        para.paragraph_format.space_after  = Pt(2)
        run = para.add_run(text)
        run.font.bold  = True
        run.font.size  = Pt(11)
        run.font.color.rgb = DARK

    def _body(text: str, italic: bool = False, indent: bool = False) -> None:
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(2)
        if indent:
            para.paragraph_format.left_indent = Cm(0.6)
        run = para.add_run(text)
        run.font.size  = Pt(10)
        run.font.italic = italic
        run.font.color.rgb = DARK

    def _bullet(text: str) -> None:
        para = doc.add_paragraph(style="List Bullet")
        para.paragraph_format.left_indent = Cm(0.8)
        para.paragraph_format.space_after = Pt(2)
        run = para.add_run(text)
        run.font.size = Pt(10)

    def _meta(text: str) -> None:
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(3)
        run = para.add_run(text)
        run.font.size  = Pt(9)
        run.font.color.rgb = GREY
        run.font.italic = True

    def _hr() -> None:
        para = doc.add_paragraph()
        para.paragraph_format.space_before = Pt(4)
        para.paragraph_format.space_after  = Pt(4)
        pPr = para._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "4")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "AAAAAA")
        pBdr.append(bottom)
        pPr.append(pBdr)

    def _monospace(text: str) -> None:
        para = doc.add_paragraph()
        run = para.add_run(text)
        run.font.name = "Courier New"
        run.font.size = Pt(9)
        run.font.color.rgb = GREY

    # ── Parse content ─────────────────────────────────────────────────────────
    for line in content.split("\n"):
        raw  = line
        line = line.strip()

        # Divider
        if re.match(r"^[━─═]{4,}$", line):
            _hr()
            continue

        if not line:
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        # Document title
        if line.startswith("LESSON PLAN:"):
            title = line.replace("LESSON PLAN:", "").strip()
            para  = doc.add_heading("", level=0)
            run   = para.add_run(title)
            run.font.bold = True
            run.font.size = Pt(18)
            run.font.color.rgb = BLUE
            continue

        # Meta line
        if line.startswith("Book:") or line.startswith("Date:"):
            _meta(line)
            continue

        # Phase header
        if re.match(
            r"^(WARM-UP|PRESENTATION|GUIDED|INDEPENDENT|ASSESSMENT|COOL-DOWN)"
            r".*\(\d+",
            line,
        ):
            _phase_heading(line)
            continue

        # Section header
        is_header = (
            len(line) >= 4
            and line == line.upper()
            and not line.startswith("•")
            and not line.startswith("|")
            and not line.startswith("-")
            and not re.match(r"^\d+\s*min", line, re.I)
        )
        if is_header:
            _heading(line, level=1)
            continue

        # Timing table rows (monospace)
        if line.startswith("|") or re.match(r"^-{3,}", line):
            _monospace(line)
            continue

        # Bullet
        if line.startswith("•"):
            _bullet(line[1:].strip())
            continue

        # Vocabulary
        if "–" in line and not line.startswith("•") and len(line) < 120:
            _body(line, italic=True, indent=True)
            continue

        # Default
        _body(line)

    # ── Footer paragraph ──────────────────────────────────────────────────────
    _hr()
    footer_para = doc.add_paragraph()
    r = footer_para.add_run("Generated by EduHelper AI — Primary Teacher Toolkit")
    r.font.size  = Pt(8)
    r.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)
    r.font.italic = True

    doc.save(str(filepath))
